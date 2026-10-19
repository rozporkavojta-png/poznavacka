import sys, json, hashlib
from pptx import Presentation
from pptx.oxml.ns import qn
res = {}
for key, path in (('p1', sys.argv[1]), ('p2', sys.argv[2])):
    prs = Presentation(path)
    out = []
    for i, s in enumerate(prs.slides, 1):
        title = ''
        for x in s.shapes:
            if x.has_text_frame and x.text_frame.text.strip():
                title = x.text_frame.text; break
        imgs = []
        for blip in s._element.iter(qn('a:blip')):
            rid = blip.get(qn('r:embed'))
            if rid and rid in s.part.rels:
                p = s.part.rels[rid].target_part
                imgs.append({'name': p.partname, 'ct': p.content_type, 'size': len(p.blob)})
        # videos / ole
        other = [x.shape_type for x in s.shapes]
        texts = [x.text_frame.text for x in s.shapes if x.has_text_frame and x.text_frame.text.strip()]
        out.append({'n': i, 'texts': texts, 'imgs': imgs})
    res[key] = out
json.dump(res, open(sys.argv[3], 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
for k in res:
    no = [s['n'] for s in res[k] if not s['imgs']]
    print(k, 'slides', len(res[k]), 'bez obrazku:', no)
