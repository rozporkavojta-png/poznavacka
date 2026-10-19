import sys, json, openpyxl
from pptx import Presentation
from pptx.util import Emu
out = {}
wb = openpyxl.load_workbook(sys.argv[1])
rows = []
for ws in wb:
    for row in ws.iter_rows():
        vals = ['' if c.value is None else str(c.value) for c in row]
        if any(vals):
            red = any(c.font and c.font.color and c.font.color.type=='rgb' and str(c.font.color.rgb).endswith('FF0000') for c in row[:2])
            rows.append({'sheet': ws.title, 'r': row[0].row, 'v': vals, 'red': red})
out['excel'] = rows
for key, path in (('p1', sys.argv[2]), ('p2', sys.argv[3])):
    prs = Presentation(path)
    slides = []
    for i, s in enumerate(prs.slides, 1):
        sh = []
        for x in s.shapes:
            d = {'type': str(x.shape_type), 'name': x.name, 'pos': [x.left, x.top, x.width, x.height]}
            if x.has_text_frame:
                d['text'] = x.text_frame.text
            if x.shape_type == 13:
                d['img'] = x.image.filename; d['ext'] = x.image.ext; d['size'] = len(x.image.blob)
            sh.append(d)
        slides.append({'n': i, 'layout': s.slide_layout.name, 'shapes': sh, 'notes': s.notes_slide.notes_text_frame.text if s.has_notes_slide else ''})
    out[key] = {'w': prs.slide_width, 'h': prs.slide_height, 'slides': slides}
json.dump(out, open(sys.argv[4], 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
