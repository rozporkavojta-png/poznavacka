import sys, json, copy, re
from pptx import Presentation
from pptx.oxml.ns import qn
from PIL import Image
sys.path.insert(0, '.')
from latin import LAT, SPLIT, MAP, MISSING

species = json.load(open('species.json', encoding='utf-8'))
zdroje = json.load(open('nove/zdroje.json', encoding='utf-8'))
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'

# opravy názvů na snímcích (první řádek) – podle Excelu
NAZEV = {
    ('p1', 110): '„skokan zelený“', ('p2', 27): 'dvojnožka dvoupórá', ('p2', 75): 'koroptev polní',
    ('p2', 103): 'dronte mauricijský (dodo)', ('p2', 138): 'buřňák lední', ('p2', 273): 'kapustňák širokonosý',
    ('p2', 274): 'dugong (moroň)', ('p2', 318): 'ksukol', ('p2', 345): 'rejsec vodní', ('p2', 391): 'zorila velká, malá',
    ('p2', 411): 'žirafa masajská', ('p2', 421): 'přímorožec jihoafrický', ('p2', 77): 'kur bankivský',
    ('p2', 191): 'ara ararauna', ('p2', 223): 'červenka obecná', ('p2', 241): 'dlask tlustozobý',
}
# překlepy v latinských názvech taxonů
PREKLEPY = {'Syngntathiformes': 'Syngnathiformes', 'Chodrichthyes': 'Chondrichthyes', 'Thinamiformes': 'Tinamiformes',
            'Ciconiifomes': 'Ciconiiformes', 'Acipitriformes': 'Accipitriformes', 'Passerifomres': 'Passeriformes',
            'Perrissodactyla': 'Perissodactyla', 'Cetartiodactila': 'Cetartiodactyla', 'Proboscidae': 'Proboscidea',
            'Macroscelidae': 'Macroscelidea', 'dlouhokžídlí': 'dlouhokřídlí', 'polopice': 'poloopice', 'matamovití': 'matamatovití', 'Psittacidae': 'Psittaciformes', 'Testudinidiae': 'Testudinidae'}

# sestavení: (část, snímek) -> (český název, latinsky)
cil = {}
for s in species:
    r = s['r']
    if r in SPLIT:
        for key, cz, lat in SPLIT[r]:
            cil[tuple(key)] = (cz, lat, r)
        continue
    if r in MISSING: continue
    key = tuple(MAP[r]) if r in MAP else tuple(s['slides'][0])
    cil[key] = (None, LAT[r], r)

def nazev_runs(p):
    """runy prvního řádku odstavce (před prvním a:br)"""
    runs = []
    for ch in p:
        if ch.tag == A + 'br': break
        if ch.tag == A + 'r': runs.append(ch)
    return runs

def nastav_titul(slide, cz, lat, key):
    title = next(x for x in slide.shapes if x.has_text_frame and x.text_frame.text.strip())
    p = title.text_frame._txBody.find(A + 'p')
    runs = nazev_runs(p)
    has_br = p.find(A + 'br') is not None
    if not has_br or key == ('p2', 27):
        # snímek bez jména druhu – vložit jméno + zalomení na začátek
        vzor = p.find(A + 'r')
        r = copy.deepcopy(vzor); rpr = r.find(A + 'rPr'); rpr.set('b', '1'); rpr.attrib.pop('i', None)
        for f in rpr.findall(A + 'solidFill'): rpr.remove(f)  # barva podle motivu jako ostatní nadpisy
        if rpr.get('sz'): rpr.set('sz', str(int(int(rpr.get('sz')) * 4 / 3)))
        r.find(A + 't').text = cz
        br = etree_br(rpr)
        idx = list(p).index(vzor)
        p.insert(idx, br); p.insert(idx, r)
        runs = [r]
    elif cz:
        runs[0].find(A + 't').text = cz
        for x in runs[1:]: p.remove(x)
        runs = runs[:1]
    else:
        # sjednotit text jména, odstranit koncové mezery
        t = ''.join(x.find(A + 't').text or '' for x in runs).strip()
        runs[0].find(A + 't').text = t
        for x in runs[1:]: p.remove(x)
        runs = runs[:1]
    # latinský název: kurzíva, bez tučného písma, 75 % velikosti
    r0 = runs[0]
    lr = copy.deepcopy(r0); rpr = lr.find(A + 'rPr')
    rpr.attrib.pop('b', None); rpr.set('i', '1'); rpr.attrib.pop('err', None)
    if rpr.get('sz'): rpr.set('sz', str(int(int(rpr.get('sz')) * 3 / 4 // 100 * 100)))
    rpr.set('lang', 'la')
    lr.find(A + 't').text = '  ' + lat
    r0.addnext(lr)
    # překlepy v ostatních runech
    for t in title.text_frame._txBody.iter(A + 't'):
        for a, b in PREKLEPY.items():
            if t.text and a in t.text: t.text = t.text.replace(a, b)

from lxml import etree
def etree_br(rpr):
    br = etree.Element(A + 'br'); br.append(copy.deepcopy(rpr)); return br

def duplikuj(prs, src, cz, lat, img_path, za_slide):
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes): sh._element.getparent().remove(sh._element)
    # oblast obrázků zdrojového snímku
    pics = [x for x in src.shapes if x._element.find('.//' + A + 'blip') is not None]
    l = min(x.left for x in pics); t = min(x.top for x in pics)
    rgt = max(x.left + x.width for x in pics); b = max(x.top + x.height for x in pics)
    for sh in src.shapes:
        if sh in pics: continue
        new.shapes._spTree.append(copy.deepcopy(sh._element))
    # pozadí
    bg = src._element.find(qn('p:cSld')).find(qn('p:bg'))
    if bg is not None:
        new._element.find(qn('p:cSld')).insert(0, copy.deepcopy(bg))
    im = Image.open(img_path); w, h = im.size
    l, t, rgt, b = 657224, 2100000, 11430000, 6650000  # jednotná plocha obsahu
    bw, bh = rgt - l, b - t
    sc = min(bw / w, bh / h); pw, ph = int(w * sc), int(h * sc)
    new.shapes.add_picture(img_path, l + (bw - pw) // 2, t + (bh - ph) // 2, pw, ph)
    nastav_titul(new, cz, lat, None)
    return new, list(prs.slides._sldIdLst)[-1]

PRED = {481}  # hyena skvrnitá před hyenu žíhanou

def zpracuj(key, path, out):
    prs = Presentation(path)
    slides = list(prs.slides)
    for (k, n), (cz, lat, r) in cil.items():
        if k != key: continue
        nastav_titul(slides[n - 1], cz if cz else NAZEV.get((k, n)), lat, (k, n))
    for (k, n), cz in NAZEV.items():
        if k == key and (k, n) not in cil: print('NAZEV bez cile', k, n)
    # nové snímky
    pridej = [(r, v) for r, v in MISSING.items() if v[0] == key]
    lst = prs.slides._sldIdLst
    ids = list(lst)
    vlozeno = {}
    for r, (k, n) in sorted(pridej):
        s = next(x for x in species if x['r'] == r)
        src = slides[n - 1]
        new, el = duplikuj(prs, src, s['cz'], LAT[r], zdroje[str(r)]['file'], n)
        # vložit za snímek n (a za dříve vložené za stejný snímek)
        lst.remove(el)
        if r in PRED:
            ids[n - 1].addprevious(el)
        else:
            anchor = vlozeno.get(n, ids[n - 1]); anchor.addnext(el); vlozeno[n] = el
        # zdroj fotky do poznámek
        new.notes_slide.notes_text_frame.text = f"Foto: Wikimedia Commons, {zdroje[str(r)]['licence']} – {zdroje[str(r)]['src']}"
    if key == 'p2':
        # dvojnožka: fotka vpravo nahoře zasahuje do nadpisu – posunout pod nadpis
        s27 = slides[26]
        for x in s27.shapes:
            if x._element.find('.//' + A + 'blip') is not None and x.top < 1900000:
                x.top = 2000000
    prs.save(out)

zpracuj('p1', sys.argv[1], sys.argv[3])
zpracuj('p2', sys.argv[2], sys.argv[4])
print('hotovo')
